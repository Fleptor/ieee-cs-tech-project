#!/usr/bin/env python3
"""Generate the finalized, 12-slide Project CIPHER Startup Alley pitch.

Install: python3 -m pip install python-pptx
Run:     python3 generate_pitch.py

The output is saved beside this script (the workspace root), independently of
the caller's working directory. No templates, images, network, or project files
are required at runtime. All artwork is editable PowerPoint geometry.

Evidence policy: architecture is grounded in the supplied repository and PDFs;
the deck makes no unverified latency or detection-accuracy promises. Physical
deployment is team-reported; integration gaps and pilot validation are explicit.
The product proposition complements an existing firewall. Speaker notes retain
sources, limitations, and the distinction between current and proposed behavior.
"""

from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise SystemExit(
        "Missing dependency. Run: python3 -m pip install python-pptx"
    ) from exc


OUTPUT = Path(__file__).resolve().parent / "CIPHER_Startup_Alley_Pitch.pptx"
WIDTH, HEIGHT = 13.333333, 7.5
BG = "081321"
PANEL = "102237"
LINE = "294359"
WHITE = "F3F7FC"
MUTED = "AFC1D3"
TEAL = "54E1C1"
BLUE = "73B9FF"
AMBER = "FFBE73"
FONT = "Noto Sans"
MONO = "JetBrains Mono"

DOC = "Documents/Project_CIPHER_full_documentation.pdf"
OUTLINE = "Documents/General Outline - The Four Horsemen.pdf"
KERNEL = "ieee-cs-tech-project/Router/ebpf/src/kernel_space.c"
DEFS = "ieee-cs-tech-project/Router/ebpf/src/shared_defs.h"
DAEMON = "ieee-cs-tech-project/Router/daemon/src/main.rs"
AI = "ieee-cs-tech-project/Server/src/ai.rs"
SERVER = "ieee-cs-tech-project/Server/src/main.rs"
HANDLERS = "ieee-cs-tech-project/Server/src/handlers.rs"


def color(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=PANEL, stroke=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if stroke:
        shape.line.color.rgb = color(stroke)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, tint=LINE, width=1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(tint)
    shape.line.width = Pt(width)
    return shape


def text(slide, value, x, y, w, h, size=24, tint=WHITE, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, shape=None):
    if shape is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shape.left, shape.top = Inches(x), Inches(y)
        shape.width, shape.height = Inches(w), Inches(h)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, row in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = row
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color(tint)
        p.space_after = Pt(0)
        p.line_spacing = 1.12
        # Override inherited title/body bullets from the standard layouts.
        p._p.get_or_add_pPr().append(OxmlElement("a:buNone"))
    return shape


def bullets(slide, items, x, y, w, size=23, step=0.8, tint=WHITE):
    """Use real editable bullet paragraphs, one deliberately short item per box."""
    for i, item in enumerate(items):
        shape = text(slide, item, x, y + i * step, w, step - 0.09, size, tint)
        p = shape.text_frame.paragraphs[0]
        props = p._p.get_or_add_pPr()
        props.remove(props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone"))
        props.set("marL", str(Inches(0.23)))
        props.set("indent", str(-Inches(0.23)))
        bullet = OxmlElement("a:buChar")
        bullet.set("char", "•")
        props.append(bullet)


def tag(slide, value, x, y, w, tint=TEAL):
    rect(slide, x, y, w, 0.36, PANEL, LINE, radius=True)
    text(slide, value, x + 0.10, y + 0.06, w - 0.20, 0.23, 10, tint, True)


def notes(slide, narrative, sources):
    slide.notes_slide.notes_text_frame.text = (
        narrative.strip() + "\n\nSOURCES / EVIDENCE\n" + "\n".join(sources)
    )


def base(prs, section, title, layout=1):
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(BG)
    title_shape = slide.shapes.title
    # Keep the actual Title placeholder; remove unused content placeholders.
    for placeholder in list(slide.placeholders):
        if placeholder != title_shape:
            placeholder._element.getparent().remove(placeholder._element)
    rect(slide, 0.55, 0.48, 0.10, 0.20, TEAL)
    text(slide, section.upper(), 0.78, 0.47, 11.7, 0.25, 11, TEAL, True)
    text(slide, title, 0.75, 1.02, 11.8, 1.30, 34, bold=True, shape=title_shape)
    line(slide, 0.75, 6.96, 12.58, 6.96)
    text(slide, "CIPHER  /  STARTUP ALLEY", 0.75, 7.12, 7, 0.20, 9, MUTED)
    text(slide, f"{len(prs.slides):02d} / 12", 11.55, 7.10, 1.03, 0.23,
         10, MUTED, align=PP_ALIGN.RIGHT, font=MONO)
    return slide


def card(slide, x, y, w, h, heading, body, accent=TEAL, number=None):
    rect(slide, x, y, w, h, PANEL, LINE, radius=True)
    rect(slide, x, y + 0.25, 0.04, h - 0.5, accent)
    if number:
        text(slide, number, x + 0.24, y + 0.25, w - 0.48, 0.60,
             30, accent, True, font=MONO)
    offset = 1.05 if number else 0.35
    text(slide, heading, x + 0.25, y + offset, w - 0.50, 0.85, 22, accent, True)
    bullets(slide, body, x + 0.25, y + offset + 1.03, w - 0.50, size=19, step=0.8)


def arrow(slide, x, y, tint=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON,
                                   Inches(x), Inches(y), Inches(0.28), Inches(0.38))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(tint)
    shape.line.fill.background()


def build_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(WIDTH), Inches(HEIGHT)
    prs.core_properties.title = "Project CIPHER | Startup Alley"
    prs.core_properties.subject = "Experimental behavioral protection alongside existing firewalls"
    prs.core_properties.author = "The 4 Horsemen"
    prs.core_properties.keywords = "CIPHER, Startup Alley, eBPF, XDP, Rust, SOHO, prototype"
    prs.core_properties.comments = "12-slide booth pitch: rehearsed software flows, explicit simulation and remaining live validation."

    nft = "https://netfilter.org/projects/nftables/manpage.html"
    suricata = "https://docs.suricata.io/en/suricata-8.0.0/capture-hardware/ebpf-xdp.html"
    opnsense = "https://docs.opnsense.org/manual/ips.html"
    faq = "CIPHER_Technical_FAQ.md"

    # 01 — A narrow, defensible product proposition.
    s = base(prs, "Network security / Prototype", "Project CIPHER", layout=0)
    text(s, "Project CIPHER", 0.75, 2.1, 8.6, 1.0, 48, bold=True, shape=s.shapes.title)
    text(s, "Behavioral protection for small networks", 0.78, 3.32, 8.4, 0.90, 27, MUTED)
    text(s, "Designed to complement an existing firewall", 0.78, 4.46, 8.4, 0.85, 20, TEAL)
    tag(s, "THE 4 HORSEMEN", 0.78, 5.57, 2.05)
    for i in range(5):
        yy = 2.40 + i * 0.60
        line(s, 9.50, yy, 10.15, yy, BLUE)
        line(s, 11.87, yy, 12.55, yy, BLUE)
        rect(s, 9.39, yy - 0.055, 0.11, 0.11, TEAL)
        rect(s, 12.55, yy - 0.055, 0.11, 0.11, TEAL)
    rect(s, 10.15, 2.04, 1.72, 3.16, PANEL, TEAL, radius=True)
    text(s, "C", 10.44, 2.78, 1.18, 1.10, 64, TEAL, True, PP_ALIGN.CENTER, MONO)
    text(s, "eBPF / XDP", 10.32, 4.27, 1.38, 0.35, 11, MUTED, align=PP_ALIGN.CENTER)
    notes(s, "CIPHER is a systems prototype exploring behavioral network protection for small offices and homes. The revised proposition is an additional analysis/enforcement layer alongside an established firewall, not a replacement for a mature stateful firewall or a proven enterprise IPS. The desired outcome is useful additional protection with manageable resource use and false blocks. No measured performance, accuracy, or commercial-readiness claim is made.", [DOC, KERNEL, DAEMON, AI, faq, "User-requested independent critique and pitch revision."])

    # 02 — Explain the gap without disparaging existing firewalls.
    s = base(prs, "02 / The problem", "Allowed traffic can still be harmful.")
    card(s, 0.75, 2.65, 5.63, 3.40, "LEGITIMATE ACTIVITY", ["A camera uploads a backup"], TEAL)
    card(s, 6.95, 2.65, 5.63, 3.40, "COMPROMISED DEVICE", ["A camera uploads stolen data"], AMBER)
    text(s, "Both may use allowed HTTPS. Context matters.", 0.80, 6.34, 11.8, 0.44, 22, MUTED)
    notes(s, "Illustrative scenario, not a successful detection claim. A firewall configured to permit a device's outbound HTTPS may allow both legitimate and malicious transfers. Metadata-based behavior analysis could add a signal, but these transfers can also look alike. The current four-feature detector and LLM are not proven to distinguish them. Do not claim that all firewalls lack advanced analytics, or that CIPHER decrypts HTTPS. The unmet need to validate is understandable, actionable device-level anomaly reporting with few disruptions.", [KERNEL, AI, nft, "Illustrative customer scenario; not a measured result."])

    # 03 — Proposed deployment, explicitly separate from current implementation.
    s = base(prs, "03 / Product approach", "Add behavioral signals to existing protection.")
    tag(s, "PROPOSED DEPLOYMENT", 0.80, 2.43, 2.82)
    card(s, 0.75, 3.04, 3.60, 2.90, "Existing firewall", ["Baseline access policy"], BLUE)
    arrow(s, 4.45, 4.14, BLUE)
    card(s, 4.86, 3.04, 3.60, 2.90, "CIPHER", ["Device-level anomalies"], TEAL)
    arrow(s, 8.57, 4.14)
    card(s, 8.98, 3.04, 3.60, 2.90, "Operator", ["Monitor, review + trust"], AMBER)
    notes(s, "The proposed topology complements an existing firewall. Monitor-and-alert policy, operator review and persistent device trust are implemented and exercised in the software demo. Monitor mode prevents NEW automatic device isolation; it does not remove existing blocks or disable kernel packet filters. Trust retains alerts and prevents automatic cloud/edge isolation; explicit manual isolation remains possible. Do not describe this as a passive packet tap. Arrows express responsibilities, not mandatory wiring. A real gateway rehearsal is still required before a pilot.", [KERNEL, DAEMON, HANDLERS, nft, "Proposed validation and deployment strategy following source review."])

    # 04 — Honest feature comparison, using an editable comparison table.
    s = base(prs, "04 / Competitive context", "Blocking is established. Detection must add value.")
    xs = [0.75, 4.35, 8.46]
    ws = [3.60, 4.11, 4.12]
    headers = ["CAPABILITY", "EXISTING TOOLS", "CIPHER TODAY"]
    rows = [
        ["Packet rules", "Standard firewall function", "Implemented"],
        ["Device anomalies", "Requires added analysis", "Experimental"],
        ["XDP acceleration", "Used by existing IPS", "Generic mode"],
        ["Comparative advantage", "Established alternatives", "Unproven"],
    ]
    for x, w, label in zip(xs, ws, headers):
        rect(s, x, 2.60, w - 0.03, 0.60, LINE)
        text(s, label, x + 0.16, 2.77, w - 0.32, 0.26, 12, TEAL, True)
    for i, row in enumerate(rows):
        yy = 3.25 + i * 0.70
        for col, (x, w, value) in enumerate(zip(xs, ws, row)):
            rect(s, x, yy, w - 0.03, 0.65, PANEL)
            text(s, value, x + 0.16, yy + 0.16, w - 0.32, 0.43,
                 17, WHITE if col < 2 else AMBER, bold=(col == 0))
    text(s, "Sources: nftables and Suricata documentation; local code review.",
         0.80, 6.49, 11.8, 0.25, 11, MUTED)
    notes(s, "Separate basic firewall capability from broader IPS products. nftables already provides kernel filtering, state tracking, address sets, counters and rate limits. Blocking ports, MACs, IPs or selected TCP flags is not the product's novelty. Basic firewall rules do not by themselves implement CIPHER's proposed per-device behavioral model. Suricata documents eBPF/XDP integration; OPNsense offers a Suricata-based IPS with Netmap. No head-to-head speed or security superiority is established. 'Implemented' means logic exists, not that all paths are validated. Commodity hardware and public libraries are implementation choices, not an exclusive moat. Defensibility would need to grow through reliable integration, deployment knowledge and demonstrated customer outcomes.", [nft, suricata, opnsense, KERNEL, DAEMON, AI])

    # 05 — Preserve technical substance without timing promises.
    s = base(prs, "05 / Technical architecture", "Keep packet handling separate from analysis.")
    card(s, 0.75, 2.65, 3.60, 3.45, "eBPF / XDP", ["Kernel enforcement"], TEAL, "DATA")
    arrow(s, 4.45, 4.03)
    card(s, 4.86, 2.65, 3.60, 3.45, "Rust daemon", ["Metadata aggregation"], BLUE, "EDGE")
    arrow(s, 8.57, 4.03, BLUE)
    card(s, 8.98, 2.65, 3.60, 3.45, "Axum server", ["Statistical analysis"], BLUE, "SERVER")
    text(s, "Generic XDP today. Throughput and latency require benchmarks.",
         0.80, 6.49, 11.8, 0.25, 12, MUTED)
    notes(s, "The data plane contains MAC/IP filtering, selected protocol checks, a Bloom-filter DNS drop path, and an 8-bit diagnostic mask. The Rust edge uses HyperLogLog for approximate unique-address counting and sends FlatBuffers summaries over WebSockets. The server uses Axum and SQLite WAL. These are useful engineering components, not guarantees of zero overhead, total memory bounds, or full attack coverage. HyperLogLog bounds one sketch, while other state still grows. FlatBuffers supports direct field reads, but outgoing messages are copied. The loader uses XdpMode::Skb; native-driver efficiency needs its own test. Earlier 300 ns and under-3-microsecond documentation claims are deliberately omitted from slides because no reproducible benchmark report was found. Cloud reasoning is not on every packet's critical path.", [KERNEL, DEFS, DAEMON, SERVER, AI, "ieee-cs-tech-project/router.fbs", DOC])

    # 06 — Make the AI an experiment, not an asserted advantage.
    s = base(prs, "06 / Detection hypothesis", "Test whether AI improves decisions.")
    card(s, 0.75, 2.65, 5.60, 3.45, "Z-score screening", ["Four telemetry features", "Saved baseline + reset"], TEAL)
    arrow(s, 6.52, 3.94)
    card(s, 6.98, 2.65, 5.60, 3.45, "LLM review", ["Baseline + recent reports", "Advisory by default"], AMBER)
    text(s, "Anthropic + compatible APIs. Offline booth uses scripted responses.",
         0.80, 6.49, 11.8, 0.25, 12, MUTED)
    notes(s, "Tier 1 learns from 50 reports: observed incoming/outgoing payload bytes, drops and port entropy. Baselines persist and can be reset. Variance floors prevent tiny changes after constant traffic from becoming extreme scores. The LLM now receives baseline means/deviations, reasons and recent report context, but no payload or reliable destination attribution. Anthropic and OpenAI Chat Completions-compatible adapters are implemented; fixture tests pass, but live credentials/model access are not yet verified. Default booth mode uses a clearly labeled deterministic fixture, NOT an AI model. Real LLM reviews can be enabled after connection testing. Explain the distinction before showing incidents. Automated cloud isolation requires enforcement mode, an untrusted device, an open incident, unchanged policy, two signals and at least 85/100 self-reported confidence. This is a policy guard, not measured accuracy. Compare against statistics alone before relying on model recommendations.", [AI, "ieee-cs-tech-project/docs/LLM_SETUP.md", "ieee-cs-tech-project/docs/BOOTH_READINESS.md"])


    # 07 — The live booth story and its boundaries.
    s = base(prs, "07 / Booth walkthrough", "Alerts stay visible. Operators stay in control.")
    tag(s, "REHEARSED SOFTWARE FLOW", 0.80, 2.44, 3.42)
    bullets(s, ["See unusual activity and its evidence", "Enable response, then trust the device", "Show provider failure and recovery"],
            0.85, 3.28, 8.5, size=25, step=0.88)
    rect(s, 10.08, 3.23, 2.50, 2.34, PANEL, LINE, radius=True)
    text(s, "DEMO", 10.31, 3.60, 2.04, 0.67, 28, AMBER, True, font=MONO)
    text(s, "simulated sensor", 10.31, 4.44, 2.04, 0.35, 13, MUTED)
    text(s, "real backend", 10.31, 5.00, 2.04, 0.35, 14, TEAL, True, font=MONO)
    notes(s, "Three-minute flow: reset; scan in monitor mode; enable enforcement then scan; trust then scan; provider outage then scan; reset. Open incident evidence, show requested versus acknowledged responses and the trusted-device exemption. The fixture labels its verdicts DEMO; no physical packets are blocked. A legitimate-download scenario changes one feature and remains allowed. Additional verified software flows include explicit command failure, reconnect/replay, false-positive review and telemetry export. The actual daemon transport also has a separate TLS reconnect/outbox test; the booth's fake sensor does not establish kernel correctness. The kernel watchdog bypasses CIPHER filtering after a stale heartbeat, but real topology and failure recovery still need hardware validation. Keep an established firewall as the baseline security boundary.", ["ieee-cs-tech-project/docs/DEMO_GUIDE.md", "ieee-cs-tech-project/scripts/booth.py", "ieee-cs-tech-project/Router/daemon/src/network.rs", KERNEL])

    # 08 — Narrow customer hypothesis without an invented market size.
    s = base(prs, "08 / Initial customer hypothesis", "Start with small offices we can support.")
    card(s, 0.75, 2.65, 5.63, 3.46, "PILOT PROFILE", ["Mixed office + IoT devices", "Compatible Linux gateway"], TEAL)
    card(s, 6.95, 2.65, 5.63, 3.46, "TEST THE NEED", ["Actionable alerts", "Manageable support burden"], BLUE)
    text(s, "Customer demand and willingness to pay remain unvalidated.",
         0.80, 6.49, 11.8, 0.25, 12, MUTED)
    notes(s, "This narrows the original SOHO ambition to manageable, engineer-supported design partners. It is a customer hypothesis, not evidence of signed customers, validated market size or willingness to pay. Commodity hardware still requires suitable kernels, permissions, NICs and deployment configuration. Test whether actionable device-level alerts and careful intervention solve a real problem beyond the site's existing firewall and DNS controls. Evaluate deployment, server/model costs, operator workload and support effort. Multi-gigabit links are one future workload to benchmark, not a claimed performance level or quantified adoption opportunity. No revenue, price, margin or growth figures are introduced.", [DOC, OUTLINE, "Proposed customer-discovery and pilot scope."])

    # 09 — A concrete comparison that can falsify the hypothesis.
    s = base(prs, "09 / Validation plan", "Prove added protection at an acceptable cost.")
    card(s, 0.75, 2.65, 3.75, 3.46, "Security", ["Missed threats", "False blocks"], TEAL)
    card(s, 4.79, 2.65, 3.75, 3.46, "Efficiency", ["CPU / RAM", "Throughput / p99 latency"], BLUE)
    card(s, 8.83, 2.65, 3.75, 3.46, "Operations", ["Installation effort", "Failure recovery"], AMBER)
    text(s, "Same hardware: firewall alone → + statistics → + LLM review",
         0.80, 6.42, 11.8, 0.44, 18, MUTED)
    notes(s, "Proposed experiment, not completed testing. Compare a well-configured baseline firewall, baseline plus CIPHER's statistical/local analysis, and baseline plus LLM review using the same hardware and traffic. Include ordinary downloads, backups, calls and updates as well as labeled malicious scenarios. First verify accurate telemetry and actual command application. Measure throughput, loss, latency distributions, CPU/RAM, false blocks per device-day, missed labeled threats, model-call cost and operator effort. Define acceptable thresholds with pilot partners before testing; do not invent targets as achieved results. Separate existing-rule drop latency from time to detect and block a new behavior. Record release, NIC/driver, kernel and XDP mode. Proceed only if additional security value outweighs disruption and operating cost; a simpler implementation may win.", [faq + ": sections 22–24.", nft, "Proposed comparative evaluation; no results claimed."])

    # 10 — Disclose the current state on the slide itself.
    s = base(prs, "10 / Prototype status", "Ready to demonstrate. Hardware proof is next.")
    card(s, 0.75, 2.65, 5.63, 3.58, "SHOW TODAY", ["Alerts, trust + incident review", "Acknowledged responses + recovery"], TEAL)
    card(s, 6.95, 2.65, 5.63, 3.58, "VALIDATE NEXT", ["Live provider account + model", "Real XDP packet blocking"], AMBER)
    text(s, "Synthetic traffic • real backend • scripted offline supervisor",
         0.80, 6.49, 11.8, 0.25, 12, MUTED)
    notes(s, "Updated after booth preparation on 17 October 2026. Builds, TLS/API integration, role checks, both supervisor adapters against fixtures, saved trust/baselines, acknowledged commands and the real daemon reconnect/outbox test pass. The presenter rehearsal checks normal traffic, single-feature download, monitor/enforce, trust, manual overrides, failed commands, provider outage/malformed response, upload anomaly, incident review, offline replay, export and reset. Browser rehearsal covers the actual dashboard. This supersedes older notes claiming missing TLS checks, authorization, database initialization and reconnect support. Live provider calls have not been made because no credentials are configured. eBPF compiles but physical packet blocking, real gateway attribution, detection accuracy and performance remain unverified. The one-interface implementation is not full bidirectional gateway attribution. Do not claim production readiness, measured accuracy, guaranteed availability, or proof of attack prevention from this simulation. See the readiness matrix and recorded demo for the exact scope.", ["ieee-cs-tech-project/docs/BOOTH_READINESS.md", "ieee-cs-tech-project/docs/DEMO_GUIDE.md", "ieee-cs-tech-project/tests/integration.py", "ieee-cs-tech-project/tests/supervisor.py", "ieee-cs-tech-project/Router/daemon/src/network.rs"])


    # 11 — Domains, without inventing credentials or one owner per card.
    s = base(prs, "11 / The team", "The 4 Horsemen")
    domains = [
        ("Networks", "Deployment + testing", TEAL),
        ("eBPF", "Packet enforcement", BLUE),
        ("Rust", "Edge + server", TEAL),
        ("AI", "Behavioral analysis", BLUE),
    ]
    for i, (domain, detail, accent) in enumerate(domains):
        card(s, 0.75 + i * 3.02, 2.68, 2.77, 3.5,
             domain, [detail], accent, f"0{i + 1}")
    notes(s, "The cards show technical domains, not four separate verified specialist credentials. The General Outline assigns systems/Rust to Zakaria, security/statistical modeling to Raneem, and full-stack/UI to Ayham; it identifies the kernel/network owner only as 'You'. Do not invent that person's name or imply separate network and eBPF headcount. The team has source spanning these domains; experienced infrastructure and security reviewers would help validate and harden the prototype.", [OUTLINE, "ieee-cs-tech-project/README.md"])

    # 12 — A credible ask tied to readiness gates.
    s = base(prs, "12 / Startup Alley ask", "Help us earn the next deployment.")
    card(s, 0.75, 2.65, 3.75, 3.55, "Design partners", ["Define SOHO pilot needs"], TEAL, "01")
    card(s, 4.79, 2.65, 3.75, 3.55, "Hardware feedback", ["Review edge/server setup"], BLUE, "02")
    card(s, 8.83, 2.65, 3.75, 3.55, "Mentors", ["Validate security + reliability"], AMBER, "03")
    text(s, "Rehearsed demo → gateway validation → measured pilots",
         0.80, 6.42, 11.8, 0.44, 19, MUTED)
    notes(s, "Ask for introductions to early SOHO design partners, feedback on the physical gateway and separate server arrangement, and infrastructure/security mentors. The sequence is intentional: demonstrate the rehearsed software flow; verify actual gateway visibility and packet enforcement; validate provider calls; agree on pilot scope and restoration; collect comparative evidence before relying on automation. Monitor-and-alert and operator review are implemented, but packet-level filtering remains active. A truly passive deployment is not established by toggling monitor mode. Keep baseline firewall protection in place. No immediate live-office installation, funding amount, signed partnership, customer count or delivery date is promised. Pilot entry and success criteria should be agreed with partners.", ["User brief: Startup Alley access, deployment feedback and mentor introductions.", faq, "Proposed readiness gates following independent critique."])

    if len(prs.slides) != 12:
        raise RuntimeError("The Startup Alley pitch must contain exactly 12 slides.")
    # Catch accidental placement outside the canvas before producing the file.
    for number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0 or
                    shape.left + shape.width > prs.slide_width or
                    shape.top + shape.height > prs.slide_height):
                raise ValueError(f"Slide {number}: shape outside slide bounds: {shape.name}")
    return prs


def main():
    prs = build_deck()
    prs.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
